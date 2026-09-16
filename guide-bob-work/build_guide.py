from pathlib import Path
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Guide_installation_prise_en_main_IBM_Bob_Work.docx"

RED = "DA1E28"
BLUE = "0F62FE"
DARK = "161616"
GRAY = "525252"
LIGHT = "F4F4F4"


def annotate(src, dst, boxes, crop=None):
    im = Image.open(src).convert("RGB")
    if crop:
        im = im.crop(crop)
        ox, oy = crop[0], crop[1]
    else:
        ox = oy = 0
    draw = ImageDraw.Draw(im)
    width = max(7, round(im.width / 280))
    for box in boxes:
        x1, y1, x2, y2 = box
        draw.rounded_rectangle((x1-ox, y1-oy, x2-ox, y2-oy), radius=14, outline=(218, 30, 40), width=width)
    im.save(dst, quality=94)


annotate(ASSETS/"api-key-create-raw.png", ASSETS/"api-key-create.png", [
    (960, 1060, 2495, 1148), (960, 1335, 2495, 1450), (2335, 1690, 2500, 1795)
], crop=(510, 280, 2920, 1850))
annotate(ASSETS/"bob-shell-raw.png", ASSETS/"bob-shell.png", [
    (1045, 635, 1425, 745), (1045, 1275, 2460, 1430)
], crop=(510, 120, 2535, 1530))
annotate(ASSETS/"runtimes-raw.png", ASSETS/"runtimes.png", [
    (1035, 600, 2455, 940), (1035, 955, 2460, 1535)
], crop=(510, 120, 2535, 1570))
annotate(ASSETS/"permissions-settings-raw.png", ASSETS/"permissions-settings.png", [
    (1035, 485, 2460, 1225), (1035, 1260, 2460, 1590)
], crop=(510, 120, 2535, 1595))
annotate(ASSETS/"new-chat-controls-raw.png", ASSETS/"new-chat-controls.png", [
    (865, 775, 1015, 885), (1705, 775, 1980, 885)
], crop=(480, 410, 2260, 915))
annotate(ASSETS/"new-chat-permissions-raw.png", ASSETS/"new-chat-permissions.png", [
    (1331, 558, 1954, 1565)
], crop=(1328, 555, 1957, 1568))


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="8"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), BLUE); rPr.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); rPr.append(underline)
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = text; r.append(t)
    h.append(r); paragraph._p.append(h)


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


def add_label_para(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + " ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK)
    p.add_run(text)
    return p


def add_figure(doc, path, caption, width=6.65):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(caption)
    c.style = doc.styles["Caption"]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)


def add_prompt(doc, n, title, purpose, prompt):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    keep_with_next(h)
    r = h.add_run(f"{n}. {title}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(DARK)
    p = doc.add_paragraph(purpose)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.05
    pPr = code._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F4F4F4"); pPr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "B8B8B8"); borders.append(e)
    pPr.append(borders)
    r = code.add_run(prompt)
    r.font.name = "Aptos Mono"
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos Mono")
    r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos Mono")
    r.font.size = Pt(9)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(DARK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for sname, size, before, after in [("Title", 28, 0, 14), ("Heading 1", 18, 14, 7), ("Heading 2", 13, 10, 4)]:
    s = styles[sname]
    s.font.name = "Aptos Display"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0,0,0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
styles["Caption"].font.name = "Aptos"
styles["Caption"].font.size = Pt(9)
styles["Caption"].font.italic = True
styles["Caption"].font.color.rgb = RGBColor.from_string(GRAY)

# Cover
p = doc.add_paragraph("IBM BOB WORK", style=None)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.runs[0]; r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(BLUE)
doc.add_paragraph("Guide d’installation et de prise en main IBM Bob Work", style="Title")
p = doc.add_paragraph("Première configuration, runtimes, permissions, voix et exemples de prompts")
p.paragraph_format.space_after = Pt(18)
p.runs[0].font.size = Pt(14); p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
tbl = doc.add_table(rows=3, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
tbl.columns[0].width = Inches(1.5); tbl.columns[1].width = Inches(5.2)
for row, (a,b) in zip(tbl.rows, [("PUBLIC", "Utilisateurs IBM Bob Work"), ("VERSION", "1.0 — 14 septembre 2026"), ("PÉRIMÈTRE", "Application Bob Work officielle — hors release Bob Work-test")]):
    row.cells[0].text=a; row.cells[1].text=b
    for c in row.cells: set_cell_border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(row.cells[0], "E8F0FE")
    row.cells[0].paragraphs[0].runs[0].bold=True
doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Important. "); r.bold=True; r.font.color.rgb=RGBColor.from_string(RED)
p.add_run("Ce guide documente exclusivement IBM Bob Work (bundle com.bobwork.desktop). Bob Work-test est une release de test et ne doit pas être utilisée pour ces étapes.")
doc.add_paragraph("Ce document accompagne l’utilisateur depuis la création d’une clé d’inférence jusqu’au lancement de premiers cas d’usage. Les cadres rouges indiquent les contrôles à utiliser. Les écrans ont été capturés sur l’application et le portail réellement installés ; aucun secret n’est affiché.")
doc.add_page_break()

doc.add_heading("Sommaire", level=1)
for item in [
    "1. Introduction et prérequis", "2. Configurer IBM Bob Shell", "3. Comprendre les runtimes et CLI",
    "4. Régler les permissions", "5. Utiliser le microphone et l’enregistrement", "6. Exemples de prompts", "7. Checklist de démarrage"
]:
    doc.add_paragraph(item, style="List Number" if False else None)
doc.add_paragraph("Les numéros de page peuvent être actualisés dans Word via Références > Table des matières si le document évolue.")

doc.add_heading("1 Introduction et prérequis", level=1)
doc.add_paragraph("Bob Work s’appuie sur une session IBM Bob existante ou sur une clé API d’inférence enregistrée dans son coffre local chiffré. Pour démarrer, vérifiez votre accès à la plateforme IBM Bob et préparez une clé dédiée à Bob Work.")
add_label_para(doc, "Objectif.", "Créer une clé limitée aux appels d’inférence et l’associer à l’équipe appropriée.")
p = add_label_para(doc, "Action.", "Ouvrez ")
add_hyperlink(p, "https://bob.ibm.com/admin/apikeys", "https://bob.ibm.com/admin/apikeys")
p.add_run(", cliquez sur Create, donnez un nom explicite à la clé, sélectionnez le scope Inference, puis l’équipe default et validez avec Create.")
add_figure(doc, ASSETS/"api-key-create.png", "Figure 1 — Création d’une clé API : scope Inference, équipe default et bouton Create.")
p = doc.add_paragraph()
r=p.add_run("Conseil de sécurité. "); r.bold=True
p.add_run("Copiez la clé au moment de sa création et conservez-la dans un emplacement sûr. Ne l’insérez jamais dans un document, une capture ou un message. Dans Bob Work, elle doit uniquement être saisie dans le champ sécurisé IBM Bob inference key.")

doc.add_heading("2 Configurer IBM Bob Shell", level=1)
add_label_para(doc, "Objectif.", "Relier Bob Work à IBM Bob et, si nécessaire, installer la version officielle de Bob Shell.")
add_label_para(doc, "Action.", "Dans Bob Work, ouvrez Settings > IBM Bob Shell. Saisissez la clé dans IBM Bob inference key, puis utilisez Save in vault. Si le shell n’est pas installé, cliquez sur Install the official version.")
add_figure(doc, ASSETS/"bob-shell.png", "Figure 2 — IBM Bob Shell : installation officielle et champ sécurisé de remplacement de la clé.")
p=doc.add_paragraph(); r=p.add_run("Libellé actuel. "); r.bold=True
p.add_run("L’interface affiche Install the official version. Ce bouton correspond à l’installation de la version officielle de Bob Shell, parfois désignée comme l’installation de la CLI IBM Shell dans des supports plus anciens.")
doc.add_paragraph("Bob Work peut réutiliser une session IBM Bob déjà ouverte. Sinon, la clé d’inférence est stockée dans le coffre local chiffré et injectée uniquement dans le processus bob run. La CLI est utile pour les scénarios en terminal et certains workflows avancés, mais elle n’est pas nécessaire à tous les usages graphiques.")

doc.add_page_break()
doc.add_heading("3 Comprendre les runtimes et CLI", level=1)
add_label_para(doc, "Objectif.", "Distinguer le socle déjà intégré des environnements optionnels installés à la demande.")
add_label_para(doc, "Action.", "Ouvrez Settings > Runtimes. Attendez la fin du chargement de l’inventaire avant d’installer, mettre à jour ou retirer un runtime.")
add_figure(doc, ASSETS/"runtimes.png", "Figure 3 — Inventaire des runtimes externes et commandes d’installation ou de mise à jour.")
doc.add_paragraph("Le socle visible comprend Bob Work core, les shared runtimes, les runtimes privés des plugins, les artifacts et le cache. Les runtimes partagés déjà intégrés et installés incluent notamment Bob Work Artifact Runtime, Diagram Runtime, Shared Python, Visualization Runtime, Excel openpyxl, LaTeX Tectonic, PDF.js, Pandoc, PowerPoint python-pptx et Word python-docx.")
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr=table.rows[0]; set_repeat_table_header(hdr)
for i,t in enumerate(["Runtime ou CLI", "Utile pour", "Nécessaire au démarrage"]):
    hdr.cells[i].text=t; set_cell_shading(hdr.cells[i], "0F62FE"); set_cell_border(hdr.cells[i]); hdr.cells[i].paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255); hdr.cells[i].paragraphs[0].runs[0].bold=True
rows=[
    ("AWS CLI v2", "Services AWS", "Non"), ("Ansible", "Automatisation d’infrastructure", "Non"),
    ("Azure CLI", "Services Azure", "Non"), ("CLI IBM Db2", "Administration Db2", "Non"),
    ("CodeGraph Runtime", "Analyse de graphes de code", "Non"), ("Docling CLI", "Conversion et extraction de documents", "Non"),
    ("Google Cloud CLI", "Services Google Cloud", "Non"), ("IBM Cloud CLI", "Services IBM Cloud", "Non"),
    ("OpenShift CLI oc", "Clusters OpenShift", "Non"), ("Qiskit Runtime", "Calcul quantique", "Non"),
    ("Terraform CLI", "Infrastructure as code", "Non"), ("Zowe CLI", "IBM Z et mainframe", "Non")]
for idx,rowdata in enumerate(rows):
    cells=table.add_row().cells
    for i,t in enumerate(rowdata):
        cells[i].text=t; set_cell_border(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
        for run in cells[i].paragraphs[0].runs: run.font.size = Pt(9)
    if idx%2: [set_cell_shading(c,"F4F7FB") for c in cells]
doc.add_paragraph("Les installations peuvent être volumineuses et prendre plusieurs minutes. Pour une utilisation standard, commencez avec les runtimes partagés déjà présents ; ajoutez une CLI uniquement lorsqu’un plugin ou une tâche en a réellement besoin.")

doc.add_heading("4 Régler les permissions", level=1)
add_label_para(doc, "Objectif.", "Choisir le niveau d’isolation et les actions que Bob peut exécuter automatiquement.")
add_label_para(doc, "Action.", "Ouvrez Settings > Permissions. L’écran observé utilise Direct disk access comme mode par défaut ; ce mode agit directement sur les fichiers du Mac selon les permissions de tâche et les règles de confidentialité macOS.")
add_figure(doc, ASSETS/"permissions-settings.png", "Figure 4 — Permissions globales et mode Direct disk access sélectionné.")
doc.add_paragraph("Les permissions globales visibles comme activées sont Read, Edit, Execute, Skill, Todo et Mode, avec Auto-approve actif. La description de l’écran mentionne aussi MCP, Subtask et Subagent ; ces options apparaissent dans le contrôle de permissions du nouveau chat.")
add_label_para(doc, "Action dans un nouveau chat.", "Cliquez sur Permissions à droite du sélecteur de projet pour ouvrir le menu de la tâche.")
add_figure(doc, ASSETS/"new-chat-permissions.png", "Figure 5 — Permissions réellement proposées dans un nouveau chat.", width=4.7)
doc.add_paragraph("La liste complète visible est : Read, Edit, Execute, MCP, Skill, Todo, Subtask, Subagent et Mode. Pour un workflow fluide, vous pouvez activer toutes les autorisations nécessaires si vous comprenez leurs effets. Limitez toutefois Edit et Execute aux espaces de travail maîtrisés, MCP aux connecteurs attendus, et Subtask ou Subagent aux travaux qui justifient une délégation.")

doc.add_page_break()
doc.add_heading("5 Utiliser le microphone et l’enregistrement", level=1)
add_label_para(doc, "Objectif.", "Dicter un prompt ou enregistrer le microphone et l’audio système depuis le composeur.")
add_label_para(doc, "Action.", "Dans un nouveau chat, utilisez l’icône de microphone pour Apple dictation et le bouton rouge pour Record microphone and system audio.")
add_figure(doc, ASSETS/"new-chat-controls.png", "Figure 6 — Contrôles de dictée, d’enregistrement et d’accès aux permissions.")
doc.add_paragraph("Au premier usage, Bob Work demande l’autorisation Microphone. La dictée nécessite également Speech Recognition. Dans Settings > Permissions, l’application affiche aussi Notifications, macOS Accessibility pour Computer Use et macOS Automation pour Chrome Control. Les boutons proposés permettent d’ouvrir les réglages système correspondants.")
p=doc.add_paragraph(); r=p.add_run("Note. "); r.bold=True
p.add_run("L’enregistrement peut inclure le microphone et l’audio système. Vérifiez le contexte de confidentialité et informez les participants avant tout enregistrement d’une réunion ou d’un appel.")

doc.add_page_break()
doc.add_heading("6 Exemples de prompts", level=1)
doc.add_paragraph("Copiez les prompts tels quels dans un nouveau chat. Chaque exemple cible une capacité précise et demande une réponse courte afin de faciliter le test.")
prompts = [
("Visualize", "Teste la création et la prévisualisation d’un graphique HTML autonome.", "@plugin:visualize Crée un fichier HTML autonome chart-test.html avec un graphique ECharts en barres (revenus Jan–Juin fictifs). Preview dans la conversation, sans Chrome. Réponds brièvement."),
("Data analytics + Visualize", "Teste une analyse simple puis sa traduction en visualisation.", "@plugin:data-analytics @plugin:visualize Analyse ces 6 mois de ventes fictives (Jan 42k, Fév 51k, Mar 48k, Avr 61k, Mai 55k, Juin 70k) : tendance, 3 insights, et génère un chart HTML analytics-test.html via Visualize (ECharts). Preview dans la conversation, sans Chrome. Réponds brièvement."),
("Cloud architect", "Teste la production d’un diagramme d’architecture Azure en SVG.", "@plugin:cloud-architect Dessine une architecture Azure simple : front web → API → AKS → Postgres. Livre un diagramme SVG professionnel dans le workspace. Réponds brièvement."),
("Designer", "Teste la maquette mobile, l’aperçu HTML et l’export Sketch.", "@plugin:ibm-agentic-designer Crée une maquette mobile d’une app de notes (liste + détail). Preview HTML dans la conversation et exporte aussi un fichier notes-app-test.sketch. Réponds brièvement."),
("PowerPoint", "Teste la génération d’une présentation structurée de cinq diapositives.", "@plugin:bob-work-microsoft-powerpoint Crée une présentation pitch-ia-test.pptx de 5 slides sur « IA agentique en entreprise » (titre, problème, solution, bénéfices, next steps). Réponds brièvement."),
("Word", "Teste la création d’un document Word court et professionnel.", "@plugin:bob-work-microsoft-word Rédige un brief d’une page brief-produit-test.docx pour une app de suivi de tâches (contexte, objectifs, utilisateurs, critères de succès). Réponds brièvement."),
("Chrome", "Teste le contrôle du navigateur et la lecture de l’onglet actif.", "@plugin:bob-work-chrome-control Ouvre https://example.com dans Chrome, confirme l’onglet actif (titre + URL), et résume en 3 puces. Réponds brièvement."),
("Plugin canvas", "Teste la création et le packaging d’un plugin sans écrasement.", "@skill:plugin-creator Crée un plugin canva-design-test-v2 qui reprend le skill de https://github.com/ComposioHQ/awesome-claude-skills/tree/master/canvas-design ainsi que les fonts du dépôt. Packaging Bob Work, sans interview longue. N’écrase pas un plugin existant. Réponds brièvement quand c’est prêt."),
("Skill file-organizer", "Teste l’import d’un skill existant dans un nouveau dossier.", "@skill:skill-creator Crée le skill file-organizer-test-v2 à partir de https://github.com/ComposioHQ/awesome-claude-skills/blob/master/file-organizer/SKILL.md. N’écrase aucun skill existant. Réponds brièvement avec le chemin."),
("Skill learning-path", "Teste la conception d’un nouveau skill avec exigences de vérification.", "@skill:skill-creator Crée le skill learning-path-test-v2 qui conçoit des parcours d’apprentissage personnalisés et progressifs (objectif, niveau, contraintes, temps, exercices, critères de validation, ressources vérifiées, jamais de liens inventés). N’écrase aucun skill existant. Réponds brièvement avec le chemin.")]
for i,(t,d,prompt) in enumerate(prompts,1): add_prompt(doc,i,t,d,prompt)

doc.add_heading("7 Checklist de démarrage", level=1)
for text in [
    "Accès à bob.ibm.com confirmé.", "Clé dédiée créée avec le scope Inference et l’équipe default.",
    "Clé stockée hors captures et enregistrée dans le coffre local de Bob Work.",
    "Version officielle de Bob Shell installée uniquement si le workflow l’exige.",
    "Runtimes optionnels installés à la demande, après lecture de leur usage.",
    "Direct disk access et permissions vérifiés avant toute action sur des fichiers réels.",
    "Microphone et Speech Recognition autorisés uniquement si la dictée ou l’enregistrement sont utilisés.",
    "Premier prompt de test exécuté dans un workspace sans données sensibles."]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(text)

# Footer with page field
for section in doc.sections:
    footer=section.footer
    p=footer.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("IBM Bob Work — Guide de prise en main   |   ")
    r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); p._p.append(fld)

doc.core_properties.title = "Guide d’installation et de prise en main IBM Bob Work"
doc.core_properties.subject = "Configuration initiale et exemples d’utilisation"
doc.core_properties.author = "IBM Bob Work"
doc.core_properties.keywords = "IBM Bob Work, Bob Shell, permissions, runtimes, API key"
doc.save(OUT)
print(OUT)
